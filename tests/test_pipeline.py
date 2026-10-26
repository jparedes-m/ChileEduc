"""Runs masterfile.py on small fake MINEDUC files (everything except scraping) and checks the outputs.

    pip install pytest openpyxl
    pytest
"""

import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
from pathlib import Path

import openpyxl
import polars as pl
import pytest

REPO = Path(__file__).resolve().parents[1]


def functions():
    namespace = {"re": re, "unicodedata": unicodedata, "pl": pl}
    exec((REPO / "code" / "0_functions.py").read_text(encoding="utf-8"), namespace)
    return namespace


@pytest.mark.parametrize("original, monthly, expected", [
    ("Rendimiento-2015.zip", False, "performance_2015"),
    ("20160202_Rendimiento_2015_PUBL.csv", False, "performance_2015"),
    ("Asistencia-declarada-mensual-Septiembre-2019.rar", True, "performance_2019_09"),
    ("Matrícula-Parvularia-2011.rar", False, "performance_2011"),
    ("Asistencia-2019.rar", True, None),
    ("manual.pdf", False, None),
])
def test_file_name_for(original, monthly, expected):
    assert functions()["file_name_for"]("performance", original, monthly) == expected


def write_zip(path, members):
    with zipfile.ZipFile(path, "w") as z:
        for name, data in members.items():
            z.writestr(name, data)


def write_mapping(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "performance"
    ws.append(["desired_names", 2015, "description"])
    for row in [["year", "AGNO", "School year"], ["rbd", "RBD", "School id | RBD"], ["cod_ense", "COD_ENSE", None],
                ["let_cur", "LET_CUR", None], ["mrun", "MRUN", None], ["prom_gral", "PROM_GRAL", "GPA"],
                ["nom_rbd", "nom_rbd", None], ["sit_fin", None, None]]:
        ws.append(row)
    ws = wb.create_sheet("assistance")
    ws.append(["desired_names", "2015_03"])
    # "ANIO" is not in the file: year must come from the file name
    for row in [["year", "ANIO"], ["month", "MES"], ["mrun", "MRUN"], ["cod_ense", "COD_ENSE"]]:
        ws.append(row)
    ws = wb.create_sheet("kinder")
    ws.append(["desired_names", 2015])
    wb.save(path)


@pytest.fixture
def project(tmp_path):
    """A copy of the repo with fake raw data, run without the scraping step."""
    shutil.copytree(REPO / "code", tmp_path / "code")
    master = (REPO / "masterfile.py").read_text(encoding="utf-8")
    (tmp_path / "masterfile.py").write_text(master.replace('source("1_scraping.py")', '# source("1_scraping.py")'))
    (tmp_path / "config").mkdir()
    write_mapping(tmp_path / "config" / "data_mapping.xlsx")

    raw = tmp_path / "data" / "raw"
    for dataset in ("performance", "assistance", "kinder"):
        (raw / dataset).mkdir(parents=True)

    # Latin-1, decimal commas, a column missing from the mapping (COD_MEN), and documentation next to the data
    perf = (
        "AGNO;RBD;COD_ENSE;LET_CUR;MRUN;PROM_GRAL;NOM_RBD;COD_MEN\n"
        "2015;100;110;A;1;6,5;Escuela Ñuñoa;1\n"
        "2015;100;110;B;2;4,1;Escuela Ñuñoa;0\n"
        "2015;100;211;A;3;5,0;Escuela Ñuñoa;0\n"  # special education: dropped
        "2015;100;310;A;4;x;Escuela Ñuñoa;0\n"    # unreadable grade: null
    ).encode("latin-1")
    write_zip(raw / "performance" / "performance_2015.zip", {
        "Rendimiento 2015/20160101_Rendimiento_2015.csv": perf,
        "Rendimiento 2015/Frecuencia_2015.csv": b"a;b\n1;2\n",
        "Rendimiento 2015/Registro_2015.pdf": b"%PDF-1.4",
    })

    # A zip inside a zip, UTF-8 with BOM
    inner = tmp_path / "inner.zip"
    write_zip(inner, {"asistencia_marzo_2015.csv": "\ufeffAGNO;MES;MRUN;COD_ENSE\n2015;3;1;110\n2015;3;2;110\n".encode()})
    write_zip(raw / "assistance" / "assistance_2015_03.zip", {"carpeta/asistencia.zip": inner.read_bytes()})

    result = subprocess.run([sys.executable, "masterfile.py"], cwd=tmp_path, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return tmp_path


def test_extraction(project):
    assert (project / "data" / "intermediate" / "performance" / "performance_2015.parquet").exists()
    docs = project / "data" / "docs" / "performance" / "performance_2015" / "Rendimiento 2015"
    assert sorted(p.name for p in docs.iterdir()) == ["Frecuencia_2015.csv", "Registro_2015.pdf"]


def test_unification(project):
    perf = pl.read_parquet(project / "data" / "clean" / "performance" / "performance_2015.parquet")
    assert perf.columns == ["year", "rbd", "cod_ense", "let_cur", "mrun", "prom_gral", "nom_rbd", "sit_fin", "cod_men"]
    assert perf["mrun"].to_list() == [1, 2, 4]
    assert perf["prom_gral"].to_list() == [6.5, 4.1, None]
    assert perf["nom_rbd"][0] == "Escuela Ñuñoa"
    assert perf.schema["year"] == pl.Int32 and perf.schema["sit_fin"] == pl.String

    asis = pl.read_parquet(project / "data" / "clean" / "assistance" / "assistance_2015_03.parquet")
    assert asis.columns == ["year", "month", "mrun", "cod_ense", "agno"]  # AGNO is unmapped, so kept as an extra
    assert asis["year"].to_list() == [2015, 2015]
    assert asis["month"].to_list() == [3, 3]


def test_dictionary(project):
    text = (project / "docs" / "dictionary" / "performance.md").read_text(encoding="utf-8")
    assert "| rbd | Int32 | School id \\| RBD | RBD | 2015 | 100.0 | 100 to 100 |" in text
    assert "| sit_fin | String |  |  |  | 0.0 |  |" in text
    assert "| COD_MEN | 2015 | 1 | cod_men |" in text


def test_sampling(project):
    sample = pl.read_parquet(project / "data" / "sample" / "performance_sample.parquet")
    assert set(sample["mrun"]) <= {1, 2, 4}


def test_rerun_keeps_existing_files(project):
    target = project / "data" / "clean" / "performance" / "performance_2015.parquet"
    before = target.stat().st_mtime_ns
    subprocess.run([sys.executable, "masterfile.py"], cwd=project, check=True, capture_output=True)
    assert target.stat().st_mtime_ns == before


def test_sample_is_by_student():
    seed, fraction = 42, 0.02
    number = ((pl.col("mrun").cast(pl.UInt64) + seed) * 2654435761 % 2**32) / 2**32
    df = pl.DataFrame({"mrun": list(range(1, 100_001)) * 2})
    picked = df.filter(number < fraction)
    students = set(picked["mrun"])
    assert 1_800 < len(students) < 2_200
    assert picked.height == 2 * len(students)  # all rows of a sampled student
    assert set(df.filter(number < 0.005)["mrun"]) <= students  # smaller samples are inside bigger ones

# [End-of-file] ----
